"""Build the Project 4 workbook: Emergency Department Wait Time Forecasting and Staffing Model."""
import os, sys, math
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, FormulaRule, DataBarRule
from lib.xl import *
from lib import cleansheets as CS
from p4_clean import run, FEATURES
from p4_ref import PROFILE, erlang_c, wq_min, req_c
from p4_milp import required_week, solve_roster

OUT = "workbook/P4_ED_Wait_Time_Forecasting_Staffing_Model.xlsx"
RUNTIME = sys.argv[1] if len(sys.argv) > 1 else "Under 1 s (measured)"
MU0 = 2.0332
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
SEASONS = ["Annual average", "Winter", "Spring", "Summer", "Fall"]
SURGES = ["None", "Flu season", "Holiday", "Flu season + holiday"]
SLAS = [15, 20, 30, 45, 60]
PROFILE_RAW = [2.4, 2.0, 1.7, 1.5, 1.4, 1.5, 2.0, 3.0, 4.3, 5.3, 5.9, 6.2, 6.1, 5.9, 5.7, 5.5, 5.4, 5.4, 5.5, 5.4, 5.0, 4.4, 3.7, 3.0]

raw, dd, cl = run()
df = cl.df
orig = list(raw.columns)
extra = ["key_repaired_flag", "patient_id_conflict_flag", "admission_date_ambiguous_flag", "discharge_date_ambiguous_flag", "outlier_fields", "outlier_count",
         "sequence_conflict_flag", "lwbs_conflict_flag", "imputed_fields", "imputed_count"] + [f[0] for f in FEATURES]
clean = df[["source_row"] + orig + extra].copy()

# ---- python reference for default scenario (roster prefill, validation)
days_rng = pd.date_range(clean.arrival_date.min(), clean.arrival_date.max())
tot = clean.arrival_date.notna().sum() / len(days_rng)
DOWIDX = [(clean.arrival_dow == nm).sum() / (days_rng.dayofweek == k).sum() / tot for k, nm in enumerate(DAYS)]
TRIAGE = clean.time_to_triage_min.mean()
TQ = 30 - TRIAGE
LAM_DEF, REQ_DEF = required_week(92000 / 365, DOWIDX, MU0, TQ, 2)
X, milp_res, A = solve_roster(REQ_DEF)
X8, X12 = X[:168], X[168:]
MILP_HOURS = int((X8 * 8).sum() + (X12 * 12).sum())

wb = Workbook(); wb.remove(wb.active)
ws_c = wb.create_sheet("Clean Data")
dump_df(ws_c, clean, fmts={c: "yyyy-mm-dd" for c in ["admission_date", "discharge_date", "arrival_date", "arrival_week_start", "arrival_month"]}, header_fill=TEAL)
ws_r = wb.create_sheet("Raw Data"); dump_df(ws_r, raw, header_fill=GREY)
ws_d = wb.create_sheet("Data Dictionary")
for r in dd.itertuples(index=False):
    ws_d.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in r])
for k, w in zip("ABCD", [34, 22, 70, 40]):
    ws_d.column_dimensions[k].width = w
for row in ws_d.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = F(9)
ws_d["A1"].font = F(14, True, NAVY)
ctx = CS.Ctx(wb, cl, raw, clean, len(raw))
for c in clean.columns:
    add_name(wb, "cd_" + c, ctx.cr(c))
N = lambda c: "cd_" + c

wl = wb.create_sheet("Lists")
lists = {"A": SEASONS, "B": SURGES, "C": SLAS, "D": DAYS}
for k, v in lists.items():
    for i, x in enumerate(v):
        wl[f"{k}{i + 2}"] = x
wl.sheet_state = "hidden"
LS = {k: f"=Lists!${k}$2:${k}${1 + len(v)}" for k, v in lists.items()}

names = ["Cover", "01 Executive Brief", "02 Hourly Staffing Optimizer", "03 Cost vs SLA Dashboard", "04 MMc Wait Calculator",
         "05 Surge Season Overlay", "06 Eight Week Backtest", "07 Solver Shift Roster", "08 Assumptions & Validation",
         "M1 Arrival Model", "M2 Erlang Staffing Table", "M3 Weekly Demand Engine"]
S = {n: wb.create_sheet(n) for n in names}
DB, M1s, M2s, M3s, RS = "'03 Cost vs SLA Dashboard'", "'M1 Arrival Model'", "'M2 Erlang Staffing Table'", "'M3 Weekly Demand Engine'", "'07 Solver Shift Roster'"

# ================================================================= dashboard selectors (global controls)
ws = S["03 Cost vs SLA Dashboard"]
setup(ws, "Labor Cost versus SLA Trade Off Dashboard", "Choose the season, surge condition, door to provider target and a day to inspect. Staffing, cost and wait respond across the whole workbook.", cols=18, width_last=9.5)
selector(ws, "B6", LS["A"], "Annual average", "SEASON", "B5"); ws.merge_cells("B6:D6")
selector(ws, "F6", LS["B"], "None", "SURGE CONDITION", "F5"); ws.merge_cells("F6:H6")
selector(ws, "J6", LS["C"], 30, "DOOR TO PROVIDER TARGET (MIN)", "J5"); ws.merge_cells("J6:L6")
selector(ws, "N6", LS["D"], "Mon", "DAY TO INSPECT", "N5"); ws.merge_cells("N6:P6")
ws.row_dimensions[6].height = 24
CT = {"season": f"{DB}!$B$6", "surge": f"{DB}!$F$6", "sla": f"{DB}!$J$6", "day": f"{DB}!$N$6"}

# ================================================================= M1 arrival model
ws = S["M1 Arrival Model"]
setup(ws, "M1  Arrival Model and Planning Inputs", "Day of week and seasonal arrival indices measured from the extract, scaled to the 92,000 annual visits in the brief, with the hourly profile and service inputs.", cols=16, width_last=11)
ws.column_dimensions["B"].width = 42
r = section(ws, 5, 2, "Planning inputs", 4)
INP = {}
items = [("annual", "Annual ED visits (brief)", 92000, "#,##0", True, "Catalogue brief: 92,000 visits a year"),
         ("sla", "Door to provider target (min)", f"={CT['sla']}", "0", False, "Selected on the dashboard"),
         ("triage", "Mean door to triage time (min), measured", f"=AVERAGE({N('time_to_triage_min')})", "0.00", False, "Clean Data"),
         ("tq", "Provider queue target = target minus triage (min)", None, "0.00", False, ""),
         ("mu", "Service rate: patients per provider hour", MU0, "0.0000", True, "Calibrated: reproduces the observed mean door to provider time at the recorded average roster (see calibration panel)"),
         ("flat", "Current roster: providers on shift every hour", f"=ROUND(AVERAGE({N('physician_on_shift_count')}),0)", "0", False, "Mean physician_on_shift_count in Clean Data"),
         ("rate", "Fully loaded cost per provider hour ($)", 180, "$#,##0", True, "Planning placeholder; replace with the local blended physician and APP rate"),
         ("floor", "Minimum providers on shift (safety floor)", 2, "0", True, "Never staff below this, whatever the demand"),
         ("flu", "Flu season arrival uplift (planning)", 1.15, "0.00", True, "Planning assumption; flu flags in the extract do not vary by season (see panel)"),
         ("hol", "Holiday arrival uplift (planning)", 1.10, "0.00", True, "Planning assumption; holiday labels in the extract contradict dates (C8)"),
         ("daily", "Base daily visits = annual / 365", None, "0.00", False, ""),
         ("sidx", "Selected season index", None, "0.000", False, ""),
         ("smult", "Selected surge multiplier", None, "0.00", False, "")]
for i, (k, lab, v, fm, is_in, cm) in enumerate(items):
    INP[k] = f"{M1s}!$D${r + i}"
for i, (k, lab, v, fm, is_in, cm) in enumerate(items):
    rr = r + i
    if k == "tq":
        v = f"={INP['sla']}-{INP['triage']}"
    if k == "daily":
        v = f"={INP['annual']}/365"
    c = kv(ws, rr, 2, lab, v, fm, is_in, span_label=2, comment=cm or None)
    if k in ("sla",):
        link_style(c, fm)
r_in_end = r + len(items)
# DOW table
r = r_in_end + 1
r = section(ws, r, 2, "Day of week index (measured)", 6)
header_row(ws, r, 2, ["Day", "Visits in extract", "Calendar days", "Visits per day", "Index"])
DOW_T = r + 1
span = f'ROW(INDIRECT(MIN({N("arrival_date")})&":"&MAX({N("arrival_date")})))'
for i, dname in enumerate(DAYS):
    rr = DOW_T + i
    vals = [dname, f'=COUNTIF({N("arrival_dow")},B{rr})', f'=SUMPRODUCT(--(WEEKDAY({span},2)={i + 1}))', f"=C{rr}/D{rr}",
            f"=E{rr}/(COUNT({N('arrival_date')})/(MAX({N('arrival_date')})-MIN({N('arrival_date')})+1))"]
    fm = [None, "#,##0", "#,##0", "0.00", "0.000"]
    for j, v in enumerate(vals):
        body_cell(ws.cell(rr, 2 + j, v), fm[j], i % 2 == 1)
DOWR = f"{M1s}!$F${DOW_T}:$F${DOW_T + 6}"
# season table
r = DOW_T + 8
r = section(ws, r, 2, "Season index (measured, meteorological seasons from arrival month)", 6)
header_row(ws, r, 2, ["Season", "Visits in extract", "Calendar days", "Visits per day", "Index"])
SEA_T = r + 1
months = {"Winter": "{12,1,2}", "Spring": "{3,4,5}", "Summer": "{6,7,8}", "Fall": "{9,10,11}"}
body_cell(ws.cell(SEA_T, 2, "Annual average")); body_cell(ws.cell(SEA_T, 6, 1), "0.000")
for i, sname in enumerate(SEASONS[1:]):
    rr = SEA_T + 1 + i
    vals = [sname, f'=COUNTIF({N("season_derived")},B{rr})', f'=SUMPRODUCT(--ISNUMBER(MATCH(MONTH({span}),{months[sname]},0)))', f"=C{rr}/D{rr}",
            f"=E{rr}/(COUNT({N('arrival_date')})/(MAX({N('arrival_date')})-MIN({N('arrival_date')})+1))"]
    fm = [None, "#,##0", "#,##0", "0.00", "0.000"]
    for j, v in enumerate(vals):
        body_cell(ws.cell(rr, 2 + j, v), fm[j], i % 2 == 0)
ws.cell(int(INP["sidx"].split("$")[-1]), 4).value = f"=INDEX($F${SEA_T}:$F${SEA_T + 4},MATCH({CT['season']},$B${SEA_T}:$B${SEA_T + 4},0))"
ws.cell(int(INP["smult"].split("$")[-1]), 4).value = f'=IF({CT["surge"]}="Flu season",{INP["flu"]},IF({CT["surge"]}="Holiday",{INP["hol"]},IF({CT["surge"]}="Flu season + holiday",{INP["flu"]}*{INP["hol"]},1)))'
# hourly profile
section(ws, 5, 8, "Hourly arrival profile (planning input)", 5)
header_row(ws, 6, 8, ["Hour", "Profile weight (input)", "Share of daily arrivals", "Arrivals per hour (base day)"])
PR0 = 7
for h in range(24):
    rr = PR0 + h
    body_cell(ws.cell(rr, 8, f"{h:02d}:00"), None, h % 2 == 1)
    c = ws.cell(rr, 9, PROFILE_RAW[h]); input_style(c, "0.0")
    body_cell(ws.cell(rr, 10, f"=I{rr}/SUM($I${PR0}:$I${PR0 + 23})"), "0.0%", h % 2 == 1)
    body_cell(ws.cell(rr, 11, f"=J{rr}*{INP['daily']}"), "0.00", h % 2 == 1)
ws.conditional_formatting.add(f"J{PR0}:J{PR0 + 23}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=0.08, color="5EA8A0"))
ws.cell(PR0 + 24, 8, "Total").font = F(9, True)
c = ws.cell(PR0 + 24, 10, f"=SUM(J{PR0}:J{PR0 + 23})"); calc_style(c, "0.0%", True)
PROF = f"{M1s}!$J${PR0}:$J${PR0 + 23}"
block_note(ws, PR0 + 26, 8, PR0 + 33, 12, "The extract records arrival dates but not arrival times, so the hour of day shape cannot be measured from it. The weights are a planning profile with a low overnight trough, a late morning peak and a sustained afternoon and evening plateau. "
           "They are normalised to 100 percent, so any values may be entered. Replace them with the hourly distribution from the ED tracking board as soon as arrival timestamps are available.")
# observed context + calibration
r = SEA_T + 7
r = section(ws, r, 2, "Observed performance and surge evidence (live)", 6)
obs = [("Observed mean door to provider (min)", f"=AVERAGE({N('door_to_provider_min')})", "0.00"),
       ("Observed share of visits within 30 minutes", f"=COUNTIF({N('sla_met_flag')},TRUE)/COUNTA({N('sla_met_flag')})", "0.0%"),
       ("Flu flag share, Winter arrivals", f'=COUNTIFS({N("season_derived")},"Winter",{N("flu_season_flag")},TRUE)/COUNTIF({N("season_derived")},"Winter")', "0.0%"),
       ("Flu flag share, Summer arrivals", f'=COUNTIFS({N("season_derived")},"Summer",{N("flu_season_flag")},TRUE)/COUNTIF({N("season_derived")},"Summer")', "0.0%"),
       ("Visits in the last 12 months of the extract", f'=COUNTIF({N("arrival_date")},">"&(MAX({N("arrival_date")})-365))', "#,##0"),
       ("Scale factor from extract to brief volume", f"={INP['annual']}/D{r + 4}", "0.0x")]
for i, (a, b, fm) in enumerate(obs):
    ws.cell(r + i, 2, a).font = F(10); c = ws.cell(r + i, 4, b); calc_style(c, fm, True, NAVY)
OBS = {"d2p": f"{M1s}!$D${r}", "sla": f"{M1s}!$D${r + 1}", "last12": f"{M1s}!$D${r + 4}"}
note(ws, r + 6, 2, "Flu flags occur at almost the same rate in Winter and Summer, so the extract cannot measure a flu surge; the uplift is therefore a planning input. The extract is a sample: its annual volume is scaled to the brief's 92,000 visits while keeping the measured day and season shape.", span=5)
r = r + 8
r = section(ws, r, 2, "Service rate calibration (annual average day, current flat roster)", 6)
header_row(ws, r, 2, ["Hour", "Arrivals per hour", "Queue wait, current roster (min, capped 240)", "Door to provider (min)"])
CAL0 = r + 1
MU, FLAT, TRI = INP["mu"], INP["flat"], INP["triage"]


def wq_f(lam, c, cap=None):
    a = f"({lam})/{MU}"
    B = f"POISSON({c},{a},FALSE)/POISSON({c},{a},TRUE)"
    core = f"IF({c}<={a},9999,{c}*{B}/({c}-{a}*(1-{B}))/({c}*{MU}-({lam}))*60)"
    return f"MIN({cap},{core})" if cap else core


for h in range(24):
    rr = CAL0 + h
    body_cell(ws.cell(rr, 2, f"{h:02d}:00"), None, h % 2 == 1)
    body_cell(ws.cell(rr, 3, f"={INP['daily']}*INDEX({PROF},{h + 1})"), "0.00", h % 2 == 1)
    body_cell(ws.cell(rr, 4, "=" + wq_f(f"C{rr}", FLAT, 240)), "0.0", h % 2 == 1)
    body_cell(ws.cell(rr, 5, f"={TRI}+D{rr}"), "0.0", h % 2 == 1)
ws.conditional_formatting.add(f"E{CAL0}:E{CAL0 + 23}", ColorScaleRule(start_type="num", start_value=10, start_color="DFF3E4", mid_type="num", mid_value=30, mid_color="FFE08A", end_type="num", end_value=90, end_color="E06666"))
rr = CAL0 + 25
cal = [("Modelled mean door to provider (arrival weighted)", f"=SUMPRODUCT(C{CAL0}:C{CAL0 + 23},E{CAL0}:E{CAL0 + 23})/SUM(C{CAL0}:C{CAL0 + 23})", "0.00"),
       ("Observed mean door to provider", f"={OBS['d2p']}", "0.00"),
       ("Calibration error (min)", f"=D{rr}-D{rr + 1}", "+0.00;-0.00"),
       ("Modelled peak hour door to provider, current roster", f"=MAX(E{CAL0}:E{CAL0 + 23})", "0.0")]
for i, (a, b, fm) in enumerate(cal):
    ws.cell(rr + i, 2, a).font = F(10, True); c = ws.cell(rr + i, 4, b); calc_style(c, fm, True, NAVY)
ws.cell(rr + 2, 5, f'=IF(ABS(D{rr + 2})<=1,"PASS","FAIL")'); CS.PASS_RULES(ws, f"E{rr + 2}")
CAL = {"model": f"{M1s}!$D${rr}", "err": f"{M1s}!$D${rr + 2}", "peak": f"{M1s}!$D${rr + 3}"}
note(ws, rr + 4, 2, "The service rate is the single free parameter of the M/M/c model. It is set so that, with arrivals shaped as above and the recorded average roster on every hour, the arrival weighted door to provider time equals the observed average. The modelled peak hour wait then serves as an independent sense check against the brief, which reports waits above 68 minutes at peaks.", span=5)
print_fit(ws, f"A1:M{rr + 6}", landscape=False, tall=2)

# ================================================================= M2 Erlang staffing table
ws = S["M2 Erlang Staffing Table"]
setup(ws, "M2  Erlang C Staffing Table", "Provider queue wait (minutes) for every arrival rate from 0.1 to 40.0 per hour and 1 to 30 providers, and the minimum providers meeting the target.", cols=34, width_last=6)
ws.column_dimensions["B"].width = 10; ws.column_dimensions["AG"].width = 11
note(ws, 4, 2, "Erlang C is computed without recursion using the identity Erlang B(c, a) = POISSON(c, a, FALSE) / POISSON(c, a, TRUE); C = c B / (c - a (1 - B)); Wq = C / (c mu - lambda). 9999 marks an unstable queue (c at or below a). "
     "Required providers = count of staffing levels whose wait exceeds the target, plus one; because Wq falls as c rises this is exactly the smallest c meeting the target.", span=33, height=30)
T0 = 8
ws.cell(T0 - 1, 2, "lambda \\ c").font = F(8, True, WHITE); ws.cell(T0 - 1, 2).fill = fill(NAVY)
for c in range(1, 31):
    h = ws.cell(T0 - 1, 2 + c, c); h.font = F(8, True, WHITE); h.fill = fill(NAVY); h.alignment = Alignment(horizontal="center")
h = ws.cell(T0 - 1, 33, "Required"); h.font = F(8, True, WHITE); h.fill = fill(TEAL)
for i in range(400):
    rr = T0 + i
    ws.cell(rr, 2, round((i + 1) / 10, 1)).font = F(8, True)
    for c in range(1, 31):
        cc = ws.cell(rr, 2 + c, "=" + wq_f(f"$B{rr}", f"{L(2 + c)}${T0 - 1}"))
        cc.font = F(7); cc.number_format = "0.0"
    cc = ws.cell(rr, 33, f'=MAX({INP["floor"]},MIN(30,COUNTIF(C{rr}:AF{rr},">"&{INP["tq"]})+1))'); cc.font = F(8, True, TEAL)
ws.conditional_formatting.add(f"C{T0}:AF{T0 + 399}", CellIsRule(operator="greaterThan", formula=[f"{INP['tq']}"], fill=fill("F6E3E3")))
ws.freeze_panes = ws.cell(T0, 3)
add_name(wb, "tbl_req", f"{M2s}!$AG${T0}:$AG${T0 + 399}")
REQ = lambda lam: f"INDEX(tbl_req,MIN(400,MAX(1,ROUNDUP(({lam})*10,0))))"

# ================================================================= 07 roster decision cells (needed by M3)
wsr = S["07 Solver Shift Roster"]
R0 = 15  # rows 15..182

# ================================================================= M3 weekly demand engine
ws = S["M3 Weekly Demand Engine"]
setup(ws, "M3  Weekly Demand and Staffing Engine (168 hours)", "Arrivals, required providers, current and optimised door to provider times and cost for every hour of the week under the dashboard scenario.", cols=20, width_last=10)
E0 = 8
hdr = ["Day", "Hour", "Hour of week", "Day index", "Hour share", "Arrivals per hour", "Required providers", "Current roster", "Door to provider, current",
       "Door to provider, optimised", "SLA met, current", "SLA met, optimised", "Shift roster coverage", "Door to provider, shift roster", "SLA met, shift roster",
       "Cost optimised ($)", "Cost current ($)", "Cost shift roster ($)"]
header_row(ws, E0 - 1, 2, hdr, height=44)
for k in range(168):
    rr = E0 + k
    d, h = divmod(k, 24)
    vals = [DAYS[d], h, k, f"=INDEX({DOWR},{d + 1})", f"=INDEX({PROF},{h + 1})",
            f"={INP['daily']}*E{rr}*F{rr}*{INP['sidx']}*{INP['smult']}", "=" + REQ(f"G{rr}"), f"={FLAT}",
            f"={TRI}+" + wq_f(f"G{rr}", f"I{rr}", 240), f"={TRI}+" + wq_f(f"G{rr}", f"H{rr}", 240),
            f"=J{rr}<={INP['sla']}", f"=K{rr}<={INP['sla']}", f"={RS}!H{R0 + k}", f"={TRI}+" + wq_f(f"G{rr}", f"N{rr}", 240), f"=O{rr}<={INP['sla']}",
            f"=H{rr}*{INP['rate']}", f"=I{rr}*{INP['rate']}", f"=N{rr}*{INP['rate']}"]
    fm = [None, "0", "0", "0.000", "0.0%", "0.00", "0", "0", "0.0", "0.0", None, None, "0", "0.0", None, "$#,##0", "$#,##0", "$#,##0"]
    for j, v in enumerate(vals):
        c = ws.cell(rr, 2 + j, v); c.font = F(8)
        if fm[j]:
            c.number_format = fm[j]
EE = E0 + 167
ws.freeze_panes = ws.cell(E0, 4)
EC = lambda col: f"{M3s}!${col}${E0}:${col}${EE}"
AGG = {"lam": EC("G"), "req": EC("H"), "cur": EC("I"), "d2c": EC("J"), "d2o": EC("K"), "okc": EC("L"), "oko": EC("M"),
       "cov": EC("N"), "d2r": EC("O"), "okr": EC("P"), "costo": EC("Q"), "costc": EC("R"), "costr": EC("S"), "day": EC("B"), "hour": EC("C")}
W = lambda col: f"SUMPRODUCT({AGG['lam']},{AGG[col]})/SUM({AGG['lam']})"

# ================================================================= 07 Solver shift roster
ws = wsr
setup(ws, "07  Solver Ready Shift Roster", "Integer shift plan (8 and 12 hour shifts, any start hour, weekly cycle) that covers the required providers at minimum paid hours. Prefilled with the optimal solution for the default scenario.", cols=14, width_last=10)
header_row(ws, R0 - 1, 2, ["Day", "Hour", "Hour of week", "8 h shifts starting (decision)", "12 h shifts starting (decision)", "Coverage", "Required",
                           "Slack", "Shortfall"], height=44)
ws.cell(R0 - 1, 8).value = "Required"
for k in range(168):
    rr = R0 + k
    d, h = divmod(k, 24)
    body_cell(ws.cell(rr, 2, DAYS[d]), None, k % 2 == 1); body_cell(ws.cell(rr, 3, h), "0", k % 2 == 1); body_cell(ws.cell(rr, 4, k), "0", k % 2 == 1)
    c = ws.cell(rr, 5, int(X8[k])); input_style(c, "0")
    c = ws.cell(rr, 6, int(X12[k])); input_style(c, "0")
    body_cell(ws.cell(rr, 7, f"=SUMPRODUCT(--(MOD($D{rr}-$D${R0}:$D${R0 + 167},168)<8),$E${R0}:$E${R0 + 167})+SUMPRODUCT(--(MOD($D{rr}-$D${R0}:$D${R0 + 167},168)<12),$F${R0}:$F${R0 + 167})"), "0", k % 2 == 1)
    body_cell(ws.cell(rr, 8, f"=INDEX({AGG['req']},{k + 1})"), "0", k % 2 == 1)
    body_cell(ws.cell(rr, 9, f"=G{rr}-H{rr}"), "+0;-0;0", k % 2 == 1)
    body_cell(ws.cell(rr, 10, f'=IF(I{rr}<0,"SHORT","")'), None, k % 2 == 1)
# coverage column in M3 refers to column H in the task spec -> coverage is column G here
for k in range(168):
    S["M3 Weekly Demand Engine"].cell(E0 + k, 14).value = f"={RS}!G{R0 + k}"
RE = R0 + 167
ws.conditional_formatting.add(f"J{R0}:J{RE}", CellIsRule(operator="equal", formula=['"SHORT"'], fill=fill(RED), font=F(9, True, WHITE)))
ws.conditional_formatting.add(f"E{R0}:F{RE}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill("FFE08A"), font=F(9, True, "0000FF")))
summ = [("Total shifts", f"=SUM(E{R0}:F{RE})", "0"), ("Paid provider hours per week", f"=8*SUM(E{R0}:E{RE})+12*SUM(F{R0}:F{RE})", "#,##0"),
        ("Weekly labor cost", f"=E6*{INP['rate']}", "$#,##0"), ("Hourly lower bound (sum of required)", f"=SUM(H{R0}:H{RE})", "#,##0"),
        ("Gap to hourly lower bound", "=E6/E8-1", "0.0%"), ("Hours with a shortfall", f'=COUNTIF(J{R0}:J{RE},"SHORT")', "0"),
        ("Current flat roster hours per week", f"={FLAT}*168", "#,##0"), ("Paid hours saved versus current", "=E11-E6", "#,##0")]
for i, (a, b, fm) in enumerate(summ):
    ws.merge_cells(start_row=5 + i, start_column=2, end_row=5 + i, end_column=4)
    ws.cell(5 + i, 2, a).font = F(10); c = ws.cell(5 + i, 5, b); calc_style(c, fm, True, NAVY)
ws.column_dimensions["B"].width = 10
ROS = {"hours": f"{RS}!$E$6", "cost": f"{RS}!$E$7", "lb": f"{RS}!$E$8", "gap": f"{RS}!$E$9", "short": f"{RS}!$E$10", "saved": f"{RS}!$E$12"}
block_note(ws, 5, 7, 10, 14, f"Optimality. The prefilled plan is the proven optimum of the integer program (HiGHS solver, status optimal): {MILP_HOURS} paid hours for the default scenario (annual average, no surge, 30 minute target). "
           f"To re solve after changing any scenario or input: Data > Solver; objective E6 (minimise); by changing E{R0}:F{R0 + 167}; constraints G{R0}:G{R0 + 167} >= H{R0}:H{R0 + 167}, E{R0}:F{R0 + 167} integer, E{R0}:F{R0 + 167} >= 0; method Simplex LP. "
           "Shortfall flags turn red whenever the current plan no longer covers the requirement.")
for k in range(5):
    ws.row_dimensions[5 + k].height = 18
ws.freeze_panes = ws.cell(R0, 5)
print_fit(ws, f"A1:N{RE}", landscape=False, tall=4)

# ================================================================= 02 Hourly staffing optimizer
ws = S["02 Hourly Staffing Optimizer"]
setup(ws, "02  Hour by Hour Staffing Optimisation", "Deliverable 1. Minimum providers meeting the door to provider target for every hour of the week, against the current flat roster, with SLA breach flags.", cols=17, width_last=8)
ws["B4"] = f'="Scenario: "&{CT["season"]}&"  |  Surge: "&{CT["surge"]}&"  |  Target: "&{CT["sla"]}&" min  |  Service rate: "&TEXT({MU},"0.00")&" patients per provider hour"'
ws["B4"].font = F(10, True, TEAL)
blocks = [("Required providers (optimised)", "H", "0", "req"), ("Door to provider with current flat roster (min)", "J", "0", "d2c"), ("Door to provider with optimised staffing (min)", "K", "0", "d2o")]
r = 6
BL = {}
for title, col, fm, key in blocks:
    section(ws, r, 2, title, 9)
    header_row(ws, r + 1, 2, ["Hour"] + DAYS + ["Week"], height=20)
    for h in range(24):
        rr = r + 2 + h
        body_cell(ws.cell(rr, 2, f"{h:02d}:00"), None, align="center")
        for d in range(7):
            body_cell(ws.cell(rr, 3 + d, f"=INDEX({EC(col)},{d * 24 + h + 1})"), fm, align="center")
        if key == "req":
            body_cell(ws.cell(rr, 10, f"=SUM(C{rr}:I{rr})"), "0", align="center")
        else:
            body_cell(ws.cell(rr, 10, f'=COUNTIF(C{rr}:I{rr},">"&{INP["sla"]})'), '0" breach"', align="center")
        ws.row_dimensions[rr].height = 13
    rng = f"C{r + 2}:I{r + 25}"
    if key == "req":
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="E8F1FB", end_type="max", end_color="2F5D8A"))
        tr = r + 26
        body_cell(ws.cell(tr, 2, "Hours"), None); ws.cell(tr, 2).font = F(9, True, NAVY)
        for d in range(8):
            c = ws.cell(tr, 3 + d, f"=SUM({L(3 + d)}{r + 2}:{L(3 + d)}{r + 25})"); calc_style(c, "0", True, NAVY); c.fill = fill(TEAL_L)
        BL["req_tot"] = f"'02 Hourly Staffing Optimizer'!$J${tr}"
    else:
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=10, start_color="DFF3E4", mid_type="num", mid_value=30, mid_color="FFE08A", end_type="num", end_value=90, end_color="E06666"))
        ws.conditional_formatting.add(f"J{r + 2}:J{r + 25}", CellIsRule(operator="greaterThan", formula=["0"], font=F(9, True, RED)))
    r += 29
# summary block to the right
section(ws, 6, 12, "Weekly comparison", 5)
cmp_ = [("Provider hours: optimised", f"=SUM({AGG['req']})", "#,##0"), ("Provider hours: current flat", f"=SUM({AGG['cur']})", "#,##0"),
        ("Provider hours: shift roster (07)", f"=SUM({AGG['cov']})", "#,##0"),
        ("Weekly cost: optimised hourly", f"=SUM({AGG['costo']})", "$#,##0"), ("Weekly cost: current", f"=SUM({AGG['costc']})", "$#,##0"),
        ("Mean door to provider: current", "=" + W("d2c"), "0.0"), ("Mean door to provider: optimised", "=" + W("d2o"), "0.0"),
        ("Hours breaching target: current", f"=COUNTIF({AGG['okc']},FALSE)", "0"), ("Hours breaching target: optimised", f"=COUNTIF({AGG['oko']},FALSE)", "0"),
        ("Peak hour wait: current (240 = unstable)", f"=MAX({AGG['d2c']})", "0.0"), ("Peak hour wait: optimised", f"=MAX({AGG['d2o']})", "0.0")]
for i, (a, b, fm) in enumerate(cmp_):
    ws.cell(7 + i, 12, a).font = F(9); c = ws.cell(7 + i, 16, b); calc_style(c, fm, True, NAVY)
ws.column_dimensions["L"].width = 12
block_note(ws, 20, 12, 30, 17, "Reading the grids. The first grid is the smallest number of providers that keeps the arrival hour's door to provider time within the target (M2 table). "
           "The second grid shows what patients experience today with the same headcount on every hour: late morning and early afternoon queues grow far beyond the target while nights are overstaffed. "
           "The third grid shows waits once staffing follows demand. Breach counts in the right hand column count days of the week breaching in that hour.")
OPTK = {"ho": f"'02 Hourly Staffing Optimizer'!$P$7", "hc": f"'02 Hourly Staffing Optimizer'!$P$8", "hr": f"'02 Hourly Staffing Optimizer'!$P$9",
        "co": f"'02 Hourly Staffing Optimizer'!$P$10", "cc": f"'02 Hourly Staffing Optimizer'!$P$11", "wc": f"'02 Hourly Staffing Optimizer'!$P$12",
        "wo": f"'02 Hourly Staffing Optimizer'!$P$13", "bc": f"'02 Hourly Staffing Optimizer'!$P$14", "bo": f"'02 Hourly Staffing Optimizer'!$P$15",
        "pc": f"'02 Hourly Staffing Optimizer'!$P$16", "po": f"'02 Hourly Staffing Optimizer'!$P$17"}
print_fit(ws, "A1:Q94", landscape=False, tall=2)

# ================================================================= 03 dashboard body
ws = S["03 Cost vs SLA Dashboard"]
cards = [("PROVIDER HOURS PER WEEK", f"={OPTK['ho']}", "#,##0", None, TEAL), ("CURRENT FLAT ROSTER HOURS", f"={OPTK['hc']}", "#,##0", "Same headcount every hour", "9AA7B8"),
         ("WEEKLY LABOR COST CHANGE", f"={OPTK['co']}-{OPTK['cc']}", '$#,##0;-$#,##0', "Optimised minus current", GREEN),
         ("MEAN DOOR TO PROVIDER", f"={OPTK['wo']}", '0.0" min"', None, TEAL), ("CURRENT MEAN WAIT", f"={OPTK['wc']}", '0.0" min"', "Flat roster, same demand", RED),
         ("HOURS BREACHING TARGET", f'={OPTK["bo"]}&" vs "&{OPTK["bc"]}', "@", "Optimised vs current, of 168", AMBER)]
for i, (a, b, fm, s_, col) in enumerate(cards):
    kpi(ws, 8, 2 + i * 3, a, b, fm, s_ or "", 3, col)
ws["B10"] = f'="Optimised; shift roster "&TEXT({OPTK["hr"]},"#,##0")'
ws["K10"] = f'="Optimised, target "&{INP["sla"]}&" min"'
for k in ("B10", "K10"):
    ws[k].font = F(8, False, GREY, True)
H0 = 30
dayoff = f"(MATCH({CT['day']},Lists!$D$2:$D$8,0)-1)*24"
ws.cell(12, H0, "Hour"); ws.cell(12, H0 + 1, "Required providers"); ws.cell(12, H0 + 2, "Current roster"); ws.cell(12, H0 + 3, "Shift roster")
ws.cell(12, H0 + 4, "Wait current"); ws.cell(12, H0 + 5, "Wait optimised"); ws.cell(12, H0 + 6, "Target"); ws.cell(12, H0 + 7, "Arrivals")
for h in range(24):
    rr = 13 + h
    ix = f"{dayoff}+{h + 1}"
    ws.cell(rr, H0, f"{h:02d}")
    ws.cell(rr, H0 + 1, f"=INDEX({AGG['req']},{ix})"); ws.cell(rr, H0 + 2, f"=INDEX({AGG['cur']},{ix})"); ws.cell(rr, H0 + 3, f"=INDEX({AGG['cov']},{ix})")
    ws.cell(rr, H0 + 4, f"=MIN(120,INDEX({AGG['d2c']},{ix}))"); ws.cell(rr, H0 + 5, f"=INDEX({AGG['d2o']},{ix})"); ws.cell(rr, H0 + 6, f"={INP['sla']}")
    ws.cell(rr, H0 + 7, f"=INDEX({AGG['lam']},{ix})")
# trade off grid for the selected day: wait for c = 1..30 at each hour, then required at each target
TG = [10, 15, 20, 25, 30, 40, 50, 60]
GC = H0 + 10
for h in range(24):
    rr = 13 + h
    for c in range(1, 31):
        ws.cell(rr, GC + c - 1, "=" + wq_f(f"${L(H0 + 7)}{rr}", c))
ws.cell(40, H0, "Target"); ws.cell(40, H0 + 1, "Provider hours (day)"); ws.cell(40, H0 + 2, "Labor cost (day)"); ws.cell(40, H0 + 3, "Mean door to provider")
for i, t in enumerate(TG):
    rr = 41 + i
    ws.cell(rr, H0, t)
    reqs = "+".join([f"MAX({INP['floor']},MIN(30,COUNTIF({L(GC)}{13 + h}:{L(GC + 29)}{13 + h},\">\"&({L(H0)}{rr}-{TRI}))+1))" for h in range(24)])
    ws.cell(rr, H0 + 1, "=" + reqs)
    ws.cell(rr, H0 + 2, f"={L(H0 + 1)}{rr}*{INP['rate']}")
for rr in range(12, 50):
    for k in range(H0, GC + 30):
        ws.cell(rr, k).font = F(7, False, GREY)
c1 = BarChart(); c1.type = "col"; ctitle(c1, "Selected day: providers needed by hour versus current roster")
c1.add_data(Reference(ws, min_col=H0 + 1, min_row=12, max_row=36), titles_from_data=True)
c1.set_categories(Reference(ws, min_col=H0, min_row=13, max_row=36))
c1.series[0].graphicalProperties.solidFill = TEAL; c1.gapWidth = 30
lc = LineChart(); lc.add_data(Reference(ws, min_col=H0 + 2, max_col=H0 + 3, min_row=12, max_row=36), titles_from_data=True)
lc.series[0].graphicalProperties.line.solidFill = RED; lc.series[0].graphicalProperties.line.dashStyle = "dash"; lc.series[0].graphicalProperties.line.width = 22000
lc.series[1].graphicalProperties.line.solidFill = NAVY; lc.series[1].graphicalProperties.line.width = 22000
c1 += lc; c1.y_axis.scaling.min = 0; c1.y_axis.title = "Providers"; c1.height = 8; c1.width = 15.5; c1.legend.position = "b"; c1.visible_cells_only = False
ws.add_chart(c1, "B12")
c2 = LineChart(); ctitle(c2, "Selected day: door to provider time by hour (min, capped at 120)")
c2.add_data(Reference(ws, min_col=H0 + 4, max_col=H0 + 6, min_row=12, max_row=36), titles_from_data=True)
c2.set_categories(Reference(ws, min_col=H0, min_row=13, max_row=36))
for s_, colr, dash in zip(c2.series, [RED, TEAL, "6B7280"], [None, None, "dash"]):
    s_.graphicalProperties.line.solidFill = colr; s_.graphicalProperties.line.width = 24000
    if dash:
        s_.graphicalProperties.line.dashStyle = dash
c2.y_axis.scaling.min = 0; c2.height = 8; c2.width = 15.5; c2.legend.position = "b"; c2.visible_cells_only = False
ws.add_chart(c2, "K12")
c3 = LineChart(); ctitle(c3, "Trade off: daily labor cost against the door to provider target")
c3.add_data(Reference(ws, min_col=H0 + 2, min_row=40, max_row=48), titles_from_data=True)
c3.set_categories(Reference(ws, min_col=H0, min_row=41, max_row=48))
c3.series[0].graphicalProperties.line.solidFill = NAVY; c3.series[0].graphicalProperties.line.width = 30000; c3.series[0].marker.symbol = "circle"
c3.y_axis.number_format = "$#,##0"; c3.x_axis.title = "Target door to provider (min)"; c3.y_axis.scaling.min = 0
c3.height = 7.5; c3.width = 15.5; c3.legend = None; c3.visible_cells_only = False
ws.add_chart(c3, "B29")
section(ws, 29, 11, "Trade off table (selected day)", 7)
header_row(ws, 30, 11, ["Target (min)", "Provider hours", "Daily labor cost", "vs 30 min target"], height=20)
for i, t in enumerate(TG):
    rr = 31 + i
    body_cell(ws.cell(rr, 11, f"={L(H0)}{41 + i}"), "0", align="center")
    body_cell(ws.cell(rr, 12, f"={L(H0 + 1)}{41 + i}"), "0", align="center")
    body_cell(ws.cell(rr, 13, f"={L(H0 + 2)}{41 + i}"), "$#,##0", align="center")
    body_cell(ws.cell(rr, 14, f"=M{rr}-$M$35"), "+$#,##0;-$#,##0;0", align="center")
ws.conditional_formatting.add("K31:N38", FormulaRule(formula=[f"$K31={INP['sla']}"], fill=fill(TEAL_L), font=F(9, True, TEAL)))
note(ws, 40, 11, "Each row re optimises every hour of the selected day for that target. Tighter targets need disproportionately more providers because queue waits fall steeply only once staffing clears the arrival rate.", span=7, height=44)
note(ws, 45, 2, "The dashed red line is today's flat roster; teal bars are the hourly optimum; the navy line is the Solver shift roster (07), which must sit on or above the bars. Waits above 120 minutes are shown at 120 for readability.", span=17)
for k in range(H0, GC + 30):
    ws.column_dimensions[L(k)].hidden = True
print_fit(ws, "A1:T47")

# ================================================================= 04 M/M/c calculator
ws = S["04 MMc Wait Calculator"]
setup(ws, "04  M/M/c Wait Time Calculator", "Deliverable 2. Queueing metrics for any arrival rate, service rate and number of providers, with a staffing curve.", cols=12, width_last=12)
ws.column_dimensions["B"].width = 44
r = section(ws, 5, 2, "Inputs", 4)
kv(ws, 6, 2, "Arrival rate lambda (patients per hour)", f"=MAX({AGG['lam']})", "0.00", False, comment="Defaults to the busiest hour of the week in the current scenario; overwrite with any value")
input_style(ws["D6"], "0.00")
kv(ws, 7, 2, "Service rate mu (patients per provider hour)", f"={MU}", "0.0000", True)
kv(ws, 8, 2, "Providers on shift c", 9, "0", True, comment="Try the current roster (8) at the busiest hour to see an unstable queue")
kv(ws, 9, 2, "Door to triage time (min)", f"={TRI}", "0.00", False); link_style(ws["D9"], "0.00")
kv(ws, 10, 2, "Door to provider target (min)", f"={INP['sla']}", "0", False); link_style(ws["D10"], "0")
lam, mu, c_ = "$D$6", "$D$7", "$D$8"
r = section(ws, 12, 2, "Results", 4)
a = f"({lam}/{mu})"
B = f"POISSON({c_},{a},FALSE)/POISSON({c_},{a},TRUE)"
res = [("Offered load a = lambda / mu (provider equivalents)", f"={a}", "0.000"),
       ("Utilisation rho = a / c", f"={a}/{c_}", "0.0%"),
       ("Queue stable (rho below 100 percent)", f'=IF(D14<1,"Yes","No: queue grows without limit")', None),
       ("Erlang B (blocking form)", f"=IF(D14<1,{B},\"\")", "0.0000"),
       ("Erlang C: probability a patient waits", f'=IF(D14<1,{c_}*D16/({c_}-{a}*(1-D16)),"")', "0.0%"),
       ("Mean wait for a provider Wq (min)", f'=IF(D14<1,D17/({c_}*{mu}-{lam})*60,"")', "0.0"),
       ("Mean door to provider (triage + Wq, min)", f'=IF(D14<1,$D$9+D18,"")', "0.0"),
       ("Mean patients waiting Lq = lambda x Wq", f'=IF(D14<1,{lam}*D18/60,"")', "0.00"),
       ("Mean time with provider 1 / mu (min)", f"=60/{mu}", "0.0"),
       ("Probability door to provider within target", f'=IF(D14<1,1-D17*EXP(-({c_}*{mu}-{lam})*MAX(0,$D$10-$D$9)/60),"")', "0.0%"),
       ("Providers needed for target (from M2 table)", f"={REQ(lam)}", "0")]
for i, (lab, f_, fm) in enumerate(res):
    ws.cell(13 + i, 2, lab).font = F(10, i in (5, 6, 9))
    c = ws.cell(13 + i, 4, f_); calc_style(c, fm, i in (5, 6, 9), NAVY if i in (5, 6, 9) else INK)
ws.conditional_formatting.add("D15", CellIsRule(operator="notEqual", formula=['"Yes"'], fill=fill(RED_L), font=F(10, True, RED)))
r = section(ws, 26, 2, "Staffing curve for this arrival rate", 8)
header_row(ws, 27, 2, ["Providers", "Utilisation", "Probability of waiting", "Wq (min)", "Door to provider (min)", "Within target probability", "Meets target"])
for i in range(20):
    rr = 28 + i
    cc = f"B{rr}"
    ws.cell(rr, 2, f"=MAX(1,CEILING({a},1))+{i - 2}")
    bb = f"POISSON({cc},{a},FALSE)/POISSON({cc},{a},TRUE)"
    vals = [f"=IF({cc}<1,\"\",{a}/{cc})", f'=IF(C{rr}>=1,"",{cc}*{bb}/({cc}-{a}*(1-{bb})))', f'=IF(C{rr}>=1,"",D{rr}/({cc}*{mu}-{lam})*60)',
            f'=IF(E{rr}="","unstable",$D$9+E{rr})', f'=IF(E{rr}="","",1-D{rr}*EXP(-({cc}*{mu}-{lam})*MAX(0,$D$10-$D$9)/60))', f'=IF(F{rr}="unstable","No",IF(F{rr}<=$D$10,"Yes","No"))']
    fm = ["0.0%", "0.0%", "0.0", "0.0", "0.0%", None]
    body_cell(ws.cell(rr, 2), "0", i % 2 == 1, align="center")
    for j, v in enumerate(vals):
        body_cell(ws.cell(rr, 3 + j, v), fm[j], i % 2 == 1, align="center")
ws.conditional_formatting.add("H28:H47", CellIsRule(operator="equal", formula=['"Yes"'], fill=fill(GREEN_L), font=F(9, True, GREEN)))
ws.conditional_formatting.add("B28:H47", FormulaRule(formula=["$B28=$D$8"], font=F(9, True, TEAL)))
ws.cell(27, 11, "Chart helper").font = F(8, False, GREY)
for i in range(20):
    rr = 28 + i
    ws.cell(rr, 11, f'=IF(ISNUMBER(F{rr}),MIN(120,F{rr}),120)').font = F(8, False, GREY)
ch = LineChart(); ctitle(ch, "Door to provider time as providers are added (min, capped at 120)")
ch.add_data(Reference(ws, min_col=11, min_row=28, max_row=47), titles_from_data=False)
ch.series[0].tx = SeriesLabel(v="Door to provider"); ch.series[0].graphicalProperties.line.solidFill = TEAL; ch.series[0].graphicalProperties.line.width = 28000
ch.set_categories(Reference(ws, min_col=2, min_row=28, max_row=47))
ch.y_axis.scaling.min = 0; ch.x_axis.title = "Providers on shift"; ch.height = 8; ch.width = 14; ch.legend = None; ch.visible_cells_only = False
ws.add_chart(ch, "F5")
ws.column_dimensions["K"].hidden = True
note(ws, 49, 2, "Model assumptions: Poisson arrivals, exponentially distributed provider time, identical providers, first come first served. The extract shows waits do not differ by acuity (08), so a single class queue matches current practice; a priority queue would lower ESI 1 and 2 waits at the expense of lower acuity patients.", span=9)
print_fit(ws, "A1:L52", landscape=False)

# ================================================================= 05 Surge overlay
ws = S["05 Surge Season Overlay"]
setup(ws, "05  Surge Season Staffing Overlay", "Deliverable 4. Weekly provider hours for every season and surge condition, and the hourly overlay for the day selected on the dashboard.", cols=14, width_last=11)
ws.column_dimensions["B"].width = 20
G0 = 40
combos = [(s, u) for s in SEASONS for u in SURGES]
ws.cell(6, G0, "base lambda")
for j, (s, u) in enumerate(combos):
    ws.cell(5, G0 + 1 + j, s); ws.cell(6, G0 + 1 + j, u)
sidx = lambda s: f"INDEX({M1s}!$F${SEA_T}:$F${SEA_T + 4},{SEASONS.index(s) + 1})"
umult = {"None": "1", "Flu season": INP["flu"], "Holiday": INP["hol"], "Flu season + holiday": f"{INP['flu']}*{INP['hol']}"}
for k in range(168):
    rr = 7 + k
    d, h = divmod(k, 24)
    ws.cell(rr, G0, f"={INP['daily']}*INDEX({DOWR},{d + 1})*INDEX({PROF},{h + 1})")
    for j, (s, u) in enumerate(combos):
        ws.cell(rr, G0 + 1 + j, "=" + REQ(f"${L(G0)}{rr}*{sidx(s)}*{umult[u]}"))
for rr in range(5, 175):
    for k in range(G0, G0 + 21):
        ws.cell(rr, k).font = F(7, False, GREY)
r = section(ws, 5, 2, "Weekly provider hours by season and surge condition", 7)
header_row(ws, r, 2, ["Season"] + SURGES + ["Flu uplift, hours"])
for i, s in enumerate(SEASONS):
    rr = r + 1 + i
    body_cell(ws.cell(rr, 2, s), None, i % 2 == 1)
    for j, u in enumerate(SURGES):
        col = L(G0 + 1 + i * 4 + j)
        body_cell(ws.cell(rr, 3 + j, f"=SUM({col}7:{col}174)"), "#,##0", i % 2 == 1, align="center")
    body_cell(ws.cell(rr, 7, f"=D{rr}-C{rr}"), "+#,##0;-#,##0;0", i % 2 == 1, align="center")
ws.conditional_formatting.add(f"C{r + 1}:F{r + 5}", ColorScaleRule(start_type="min", start_color="E8F1FB", end_type="max", end_color="2F5D8A"))
SURT = r + 1
r = SURT + 6
r = section(ws, r, 2, "Weekly labor cost by season and surge condition", 7)
header_row(ws, r, 2, ["Season"] + SURGES + ["Flu uplift, cost"])
for i, s in enumerate(SEASONS):
    rr = r + 1 + i
    body_cell(ws.cell(rr, 2, s), None, i % 2 == 1)
    for j in range(4):
        body_cell(ws.cell(rr, 3 + j, f"={L(3 + j)}{SURT + i}*{INP['rate']}"), "$#,##0", i % 2 == 1, align="center")
    body_cell(ws.cell(rr, 7, f"=D{rr}-C{rr}"), "+$#,##0;-$#,##0;0", i % 2 == 1, align="center")
r = r + 7
ws.cell(r, 2, "Extra weekly provider hours in the worst case (flu season + holiday, highest season) versus annual average with no surge").font = F(10, True)
c = ws.cell(r + 1, 2, f"=MAX(F{SURT}:F{SURT + 4})-C{SURT}"); calc_style(c, '+#,##0" provider hours"', True, RED)
SURGE_MAX = f"'05 Surge Season Overlay'!$B${r + 1}"
# overlay for selected day and season
H1 = G0 + 22
dayoff = f"(MATCH({CT['day']},Lists!$D$2:$D$8,0)-1)*24"
ws.cell(6, H1, "Hour")
for j, u in enumerate(SURGES):
    ws.cell(6, H1 + 1 + j, u)
for h in range(24):
    rr = 7 + h
    ws.cell(rr, H1, f"{h:02d}")
    for j, u in enumerate(SURGES):
        ws.cell(rr, H1 + 1 + j, f"=INDEX({L(G0 + 1 + j)}7:{L(G0 + 20)}174,{dayoff}+{h + 1},(MATCH({CT['season']},Lists!$A$2:$A$6,0)-1)*4+1)")
for rr in range(6, 31):
    for k in range(H1, H1 + 5):
        ws.cell(rr, k).font = F(7, False, GREY)
ch = LineChart(); ctitle(ch, "Selected day and season: providers needed under each surge condition")
ch.add_data(Reference(ws, min_col=H1 + 1, max_col=H1 + 4, min_row=6, max_row=30), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=H1, min_row=7, max_row=30))
for s_, colr, wdt in zip(ch.series, ["9AA7B8", AMBER, "6B7280", RED], [28000, 22000, 18000, 22000]):
    s_.graphicalProperties.line.solidFill = colr; s_.graphicalProperties.line.width = wdt
ch.y_axis.scaling.min = 0; ch.y_axis.title = "Providers"; ch.height = 8.5; ch.width = 16; ch.legend.position = "b"; ch.visible_cells_only = False
ws.add_chart(ch, "I5")
note(ws, r + 3, 2, "Surge multipliers are planning inputs on M1 (flu uplift and holiday uplift) because the extract's flu and holiday labels do not track volume. The overlay recomputes every hour of the week from the M2 staffing table for each of the 20 season and surge combinations.", span=12)
for k in range(G0, H1 + 5):
    ws.column_dimensions[L(k)].hidden = True
print_fit(ws, f"A1:P{r + 5}", landscape=False)

# ================================================================= 06 Backtest
ws = S["06 Eight Week Backtest"]
setup(ws, "06  Eight Week Historical Peak Backtest", "Deliverable 5. The eight highest volume weeks in the extract are forecast, staffed with the model's rule and exposed to the volume that arrived, under three lenses.", cols=18, width_last=9.5)
ws.column_dimensions["B"].width = 13
wk = clean.dropna(subset=["arrival_date"]).groupby("arrival_week_start").size()
first, last = clean.arrival_date.min(), clean.arrival_date.max()
wk = wk[(wk.index >= first) & (wk.index + pd.Timedelta(days=6) <= last)]
peak_weeks = sorted(wk.sort_values(ascending=False).index[:8])
all_weeks = list(wk.index)
kv(ws, 5, 2, "Pass rule: share of hours meeting target at least", 0.90, "0%", True, span_label=5)
kv(ws, 6, 2, "Pass rule: mean door to provider at most (min)", 35, "0", True, span_label=5)
kv(ws, 7, 2, "Surge buffer used for the buffered roster (flu x holiday on M1)", f"={INP['flu']}*{INP['hol']}", "0.000", False, span_label=5)
# weekly series for noise diagnostics (hidden block)
WS0 = 12
ws.cell(WS0 - 1, 80, "Complete weeks"); ws.cell(WS0 - 1, 81, "Visits")
for i, w in enumerate(all_weeks):
    ws.cell(WS0 + i, 80, w.to_pydatetime()); ws.cell(WS0 + i, 81, f"=COUNTIFS({N('arrival_week_start')},{L(80)}{WS0 + i})")
WR = f"$CC${WS0}:$CC${WS0 + len(all_weeks) - 1}"
section(ws, 4, 10, "Is the weekly variation real? (live)", 8)
diag = [("Complete weeks in extract", f"=COUNT({WR})", "0"), ("Mean visits per week (extract scale)", f"=AVERAGE({WR})", "0.0"),
        ("Variance of weekly visits", f"=VAR({WR})", "0.0"), ("Dispersion index (variance / mean; 1 = pure chance)", "=Q7/Q6", "0.00"),
        ("Reliability: share of variation that is real signal", "=MAX(0,1-Q6/Q7)", "0.0%"),
        ("Expected busiest of these weeks by chance alone", f"=Q6+NORMSINV(1-1/Q5)*SQRT(Q6)", "0.0")]
for i, (a_, b_, fm) in enumerate(diag):
    ws.cell(5 + i, 10, a_).font = F(9); c = ws.cell(5 + i, 17, b_); calc_style(c, fm, True, NAVY)
REL = "$Q$9"
hdr = ["Week starting", "Season", "Visits forecast", "Visits observed", "Raw uplift", "Signal adjusted uplift",
       "Raw: hours met", "Raw: mean wait", "Raw", "Adjusted: hours met", "Adjusted: mean wait", "Adjusted",
       "Buffered roster: hours met", "Buffered roster: mean wait", "Buffered", "Observed wait in extract (min)"]
header_row(ws, 12, 2, hdr, height=52)
BG = 30
PERW = 5
for w_i, w in enumerate(peak_weeks):
    rr = 13 + w_i
    season = {12: "Winter", 1: "Winter", 2: "Winter", 3: "Spring", 4: "Spring", 5: "Spring", 6: "Summer", 7: "Summer", 8: "Summer", 9: "Fall", 10: "Fall", 11: "Fall"}[(w + pd.Timedelta(days=3)).month]
    cq, cb, dr, da, db = [L(BG + w_i * PERW + k) for k in range(PERW)]
    sidx_ = f"INDEX({M1s}!$F${SEA_T}:$F${SEA_T + 4},MATCH($C${rr},{M1s}!$B${SEA_T}:$B${SEA_T + 4},0))"
    wsum = lambda col: f"SUMPRODUCT({L(BG + 8 * PERW)}$20:{L(BG + 8 * PERW)}$187,{col}$20:{col}$187)/SUM({L(BG + 8 * PERW)}$20:{L(BG + 8 * PERW)}$187)"
    vals = [w.to_pydatetime(), season,
            f"=COUNT({N('arrival_date')})/(MAX({N('arrival_date')})-MIN({N('arrival_date')})+1)*SUM({DOWR})*{sidx_}",
            f"=COUNTIFS({N('arrival_week_start')},B{rr})", f"=E{rr}/D{rr}", f"=1+(F{rr}-1)*{REL}",
            f'=COUNTIF({dr}$20:{dr}$187,"<="&{INP["sla"]})/168', "=" + wsum(dr), f'=IF(AND(H{rr}>=$G$5,I{rr}<=$G$6),"PASS","FAIL")',
            f'=COUNTIF({da}$20:{da}$187,"<="&{INP["sla"]})/168', "=" + wsum(da), f'=IF(AND(K{rr}>=$G$5,L{rr}<=$G$6),"PASS","FAIL")',
            f'=COUNTIF({db}$20:{db}$187,"<="&{INP["sla"]})/168', "=" + wsum(db), f'=IF(AND(N{rr}>=$G$5,O{rr}<=$G$6),"PASS","FAIL")',
            f"=AVERAGEIFS({N('door_to_provider_min')},{N('arrival_week_start')},B{rr})"]
    fm = ["dd mmm yy", None, "0.0", "0", "0.00x", "0.000x", "0%", "0.0", None, "0%", "0.0", None, "0%", "0.0", None, "0.0"]
    for j, v in enumerate(vals):
        body_cell(ws.cell(rr, 2 + j, v), fm[j], w_i % 2 == 1, align="center")
    for k in range(168):
        gr = 20 + k
        base = f"${L(BG + 8 * PERW)}{gr}*{sidx_}"
        ws.cell(gr, BG + w_i * PERW, "=" + REQ(base))
        ws.cell(gr, BG + w_i * PERW + 1, "=" + REQ(f"{base}*$G$7"))
        ws.cell(gr, BG + w_i * PERW + 2, f"={TRI}+" + wq_f(f"({base}*$F${rr})", f"{cq}{gr}", 240))
        ws.cell(gr, BG + w_i * PERW + 3, f"={TRI}+" + wq_f(f"({base}*$G${rr})", f"{cq}{gr}", 240))
        ws.cell(gr, BG + w_i * PERW + 4, f"={TRI}+" + wq_f(f"({base}*$F${rr})", f"{cb}{gr}", 240))
for k in range(168):
    d, h = divmod(k, 24)
    ws.cell(20 + k, BG + 8 * PERW, f"={INP['daily']}*INDEX({DOWR},{d + 1})*INDEX({PROF},{h + 1})")
for rr in range(11, 190):
    for k in list(range(BG, BG + 8 * PERW + 1)) + [80, 81]:
        ws.cell(rr, k).font = F(7, False, GREY)
CS.PASS_RULES(ws, "J13:J20"); CS.PASS_RULES(ws, "M13:M20"); CS.PASS_RULES(ws, "P13:P20")
ws.conditional_formatting.add("F13:F20", DataBarRule(start_type="num", start_value=1, end_type="num", end_value=1.5, color="E29578"))
summ = [("Weeks backtested", "=COUNT(E13:E20)", '0" of 8"'), ("Pass: raw uplift, forecast roster", '=COUNTIF(J13:J20,"PASS")', '0" of 8"'),
        ("Pass: signal adjusted uplift, forecast roster", '=COUNTIF(M13:M20,"PASS")', '0" of 8"'),
        ("Pass: raw uplift, surge buffered roster", '=COUNTIF(P13:P20,"PASS")', '0" of 8"'),
        ("Mean absolute weekly forecast error (raw)", "=SUMPRODUCT(ABS(F13:F20-1))/8", "0.0%")]
for i, (a_, b_, fm) in enumerate(summ):
    ws.cell(23 + i, 2, a_).font = F(10, True); c = ws.cell(23 + i, 6, b_); calc_style(c, fm, True, NAVY)
BT = {"n": "'06 Eight Week Backtest'!$F$23", "raw": "'06 Eight Week Backtest'!$F$24", "adj": "'06 Eight Week Backtest'!$F$25", "buf": "'06 Eight Week Backtest'!$F$26", "pass": "'06 Eight Week Backtest'!$F$25"}
ch = BarChart(); ch.type = "col"; ctitle(ch, "Peak weeks: forecast versus observed visits (extract scale)")
ch.add_data(Reference(ws, min_col=4, max_col=5, min_row=12, max_row=20), titles_from_data=True)
for i in range(8):
    c = ws.cell(13 + i, 79, f'=TEXT(B{13 + i},"dd mmm yy")'); c.font = F(7, False, GREY)
ws.column_dimensions[L(79)].hidden = True
ch.set_categories(Reference(ws, min_col=79, min_row=13, max_row=20))
ch.series[0].graphicalProperties.solidFill = "B8C4D6"; ch.series[1].graphicalProperties.solidFill = NAVY
ch.y_axis.scaling.min = 0; ch.height = 7.5; ch.width = 15; ch.legend.position = "b"; ch.gapWidth = 60; ch.visible_cells_only = False
ws.add_chart(ch, "B29")
block_note(ws, 23, 10, 44, 18, "How to read the three lenses. Peak weeks are the eight complete weeks with the most arrivals. Each is forecast from the extract's average volume and calendar indices, staffed hour by hour with the M2 rule, and then exposed to demand. "
           "Raw applies the week's observed uplift to full scale demand. But the extract holds only about 72 visits a week, so its weeks vary by chance alone by roughly 12 percent; the diagnostics above show that almost all of the week to week variation is chance (dispersion index close to 1), and the busiest weeks sit where chance alone would put them. "
           "Picking the eight busiest of 208 such weeks therefore selects sampling noise, and applying that noise to 92,000 visits a year overstates real volatility. "
           "Signal adjusted shrinks each uplift by the measured reliability (empirical Bayes), the statistically correct estimate of the real surge. "
           "Buffered roster keeps the raw, inflated uplift but staffs the week at the flu and holiday planning level, testing whether the recommended surge policy would have absorbed even the exaggerated peaks. "
           "A week passes if at least 90 percent of hours meet the target and the arrival weighted mean door to provider is within 35 minutes.")
for k in list(range(BG, BG + 8 * PERW + 1)) + [80, 81]:
    ws.column_dimensions[L(k)].hidden = True
print_fit(ws, "A1:S46")

# ================================================================= 08 Assumptions & validation
ws = S["08 Assumptions & Validation"]
setup(ws, "08  Assumptions, Validation and Limitations", "Queueing formulas checked against an independent Python implementation, calibration evidence, success metric log and assumption register.", cols=10, width_last=12)
ws.column_dimensions["B"].width = 30
r = section(ws, 5, 2, "V1  Erlang C wait: Excel POISSON identity versus Python recursion", 9)
cases = [(5, 2.0, 3), (10, 2.0, 6), (15.9, 2.0332, 8), (15.9, 2.0332, 9), (8, 1.5, 6), (20, 2.5, 9), (3, 2.0, 2), (25, 2.0, 14), (12.5, 2.2, 7), (1.2, 2.0, 2), (30, 2.0, 18), (18, 2.0332, 12)]
header_row(ws, r, 2, ["Case (lambda, mu, c)", "Python Wq (min)", "Excel Wq (min)", "Difference", "Check"])
V1 = r + 1
for i, (lm, m, cc) in enumerate(cases):
    rr = V1 + i
    ws.cell(rr, 9, lm); ws.cell(rr, 10, m); ws.cell(rr, 11, cc)
    a_ = f"(I{rr}/J{rr})"; B_ = f"POISSON(K{rr},{a_},FALSE)/POISSON(K{rr},{a_},TRUE)"
    ex = f"=K{rr}*{B_}/(K{rr}-{a_}*(1-{B_}))/(K{rr}*J{rr}-I{rr})*60"
    vals = [f"lambda {lm}, mu {m}, c {cc}", round(wq_min(lm, m, cc), 10), ex, f"=D{rr}/C{rr}-1", f'=IF(ABS(E{rr})<0.0001,"PASS","FAIL")']
    fm = [None, "0.0000", "0.0000", "+0.00000%;-0.00000%", None]
    for j, v in enumerate(vals):
        body_cell(ws.cell(rr, 2 + j, v), fm[j], i % 2 == 1)
    for k in (9, 10, 11):
        ws.cell(rr, k).font = F(7, False, GREY)
V1E = V1 + len(cases) - 1
CS.PASS_RULES(ws, f"F{V1}:F{V1E}")
for k in (9, 10, 11):
    ws.column_dimensions[L(k)].hidden = True
r = V1E + 2
r = section(ws, r, 2, "V2  Required providers, 168 hours, default scenario: Excel engine versus Python reference", 9)
ws.cell(r, 2, "Hours where Excel and Python agree").font = F(10)
ref_str = ",".join(str(x) for x in REQ_DEF)
ws.cell(r, 9, ref_str).font = F(6, False, WHITE)
header_row(ws, r + 1, 2, ["Hour of week", "Python required", "Excel required", "Match"])
V2 = r + 2
for k in range(168):
    rr = V2 + k
    ws.cell(rr, 2, k); ws.cell(rr, 3, REQ_DEF[k]); ws.cell(rr, 4, f"=INDEX({AGG['req']},{k + 1})"); ws.cell(rr, 5, f"=C{rr}=D{rr}")
    for j in range(2, 6):
        ws.cell(rr, j).font = F(7, False, GREY)
    ws.row_dimensions[rr].hidden = True
c = ws.cell(r, 5, f"=COUNTIF(E{V2}:E{V2 + 167},TRUE)"); calc_style(c, '0" of 168"', True, NAVY)
V2RES = f"'08 Assumptions & Validation'!$E${r}"
r = V2 + 169
r = section(ws, r, 2, "Success metric log", 9)
header_row(ws, r, 2, ["Metric (catalogue)", "Target", "Result", "Status", "Evidence"])
ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=10)
mets = [("Projected average wait time", "under 35 min", f"={OPTK['wo']}", '0.0" min"', f'=IF(D{{r}}<35,"MET","NOT MET")', "Arrival weighted door to provider, optimised staffing, dashboard scenario (02)."),
        ("Labor cost within 3% of Solver optimum", "within 3%", f"=({ROS['hours']}-{MILP_HOURS})/{MILP_HOURS}", "0.0%", f'=IF(ABS(D{{r}})<=0.03,"MET","NOT MET")', f"Shift roster (07) versus the proven integer optimum of {MILP_HOURS} paid hours; gap to the unrosterable hourly lower bound is shown on 07."),
        ("All eight historical peak weeks backtested", "8 of 8", f"={BT['adj']}", '0" of 8"', f'=IF(D{{r}}=8,"MET","NOT MET")', f'="06: all 8 weeks run. Signal adjusted lens shown; raw sample uplift passes "&{BT["raw"]}&" of 8 and the surge buffered roster "&{BT["buf"]}&" of 8."'),
        ("SLA breach flags accurate within one hour block", "1 hour", "Hourly by design", None, '="NOT TESTABLE"', "Flags are computed per hour (02, M3). Accuracy against observed hourly breaches cannot be measured because the extract has no arrival or provider timestamps."),
        ("Erlang C implementation correct", "exact", f'=COUNTIF(F{V1}:F{V1E},"PASS")&" of 12"', None, f'=IF(COUNTIF(F{V1}:F{V1E},"PASS")=12,"PASS","FAIL")', "V1 above."),
        ("Service rate calibration", "within 1 min", f"={CAL['err']}", '+0.00" min";-0.00" min"', f'=IF(ABS(D{{r}})<=1,"PASS","FAIL")', "M1 calibration panel."),
        ("Full recalculation responsiveness", "fast", RUNTIME, None, '="MEASURED"', "Scenario selector change timed in LibreOffice headless, single thread."),
        ("Cleaning row reconciliation", "0", "C11", "#,##0", f'=IF(D{{r}}=0,"PASS","FAIL")', "C11 audit log.")]
MT0 = r + 1
for i, (a_, b_, c_, fm, st, ev) in enumerate(mets):
    rr = MT0 + i
    ws.merge_cells(start_row=rr, start_column=6, end_row=rr, end_column=10)
    for k, v, f_ in ((2, a_, None), (3, b_, None), (4, c_, fm), (5, st.format(r=rr), None), (6, ev, None)):
        body_cell(ws.cell(rr, k, v), f_, i % 2 == 1, align="wrap" if k in (2, 6) else "center")
    ws.row_dimensions[rr].height = 34
    if a_.startswith("Cleaning"):
        REC_ROW = rr
for lab_, col_, fc_ in (("MET", GREEN_L, GREEN), ("PASS", GREEN_L, GREEN), ("NOT MET", RED_L, RED), ("FAIL", RED_L, RED), ("NOT TESTABLE", AMBER_L, AMBER), ("MEASURED", GREEN_L, GREEN)):
    ws.conditional_formatting.add(f"E{MT0}:E{MT0 + 7}", CellIsRule(operator="equal", formula=[f'"{lab_}"'], fill=fill(col_), font=F(9, True, fc_)))
METS = {r_: f"'08 Assumptions & Validation'!$E${MT0 + i}" for i, r_ in enumerate(["wait", "cost", "bt", "flags"])}
METR = {r_: f"'08 Assumptions & Validation'!$D${MT0 + i}" for i, r_ in enumerate(["wait", "cost", "bt", "flags"])}
r = MT0 + 9
r = section(ws, r, 2, "Assumption register", 9)
asm = [("Annual volume", "92,000 visits (brief); extract scaled by the factor on M1 while keeping measured day and season shape."),
       ("Hourly profile", "Planning input (M1); the extract has no arrival times."),
       ("Service rate", f"{MU0} patients per provider hour, calibrated to the observed 37.0 minute mean door to provider at the recorded average roster of 8."),
       ("Queue model", "M/M/c, first come first served, providers identical; consistent with waits that do not vary by ESI level in the extract."),
       ("Door to provider", "Triage time (measured mean) plus provider queue wait; the target applies to the sum."),
       ("Surge multipliers", "Flu 1.15 and holiday 1.10 are planning inputs; extract flags do not track volume."),
       ("Cost", "180 dollars per provider hour is a placeholder; results scale linearly with it."),
       ("Shift rules", "8 and 12 hour shifts, any start hour, weekly cycle; breaks and handover overlap not modelled.")]
for i, (a_, b_) in enumerate(asm):
    rr = r + i
    body_cell(ws.cell(rr, 2, a_), None, i % 2 == 1)
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=10)
    body_cell(ws.cell(rr, 3, b_), None, i % 2 == 1, align="wrap")
    ws.row_dimensions[rr].height = 26
r += len(asm) + 1
r = section(ws, r, 2, "Findings from the extract", 9)
fnd = ["Waits do not vary by acuity: mean door to provider is about 37 minutes for every ESI level, including ESI 1, and 383 ESI 1 patients waited more than 10 minutes. This is a clinical safety issue independent of staffing levels.",
       "Only 39.4 percent of visits met the 30 minute target in the extract.",
       "Day of week and seasonal volume vary by less than 3 percent in this extract, so the staffing pattern is driven almost entirely by the hour of day."]
for t in fnd:
    r = note(ws, r, 2, t, span=9, color=INK)
print_fit(ws, f"A1:K{r + 1}", landscape=False, tall=2)

# ================================================================= 01 Executive brief
ws = S["01 Executive Brief"]
setup(ws, "Staffing Recommendation: Executive Brief", "Emergency department provider staffing to meet the 30 minute door to provider target at 92,000 annual visits. Figures follow the dashboard scenario.", cols=12, width_last=10.5)
kp = [("MEAN DOOR TO PROVIDER, TODAY", f"={OPTK['wc']}", '0.0" min"', "Flat roster, modelled", RED), ("MEAN DOOR TO PROVIDER, PLAN", f"={OPTK['wo']}", '0.0" min"', "Demand based staffing", TEAL),
      ("PEAK HOUR WAIT, TODAY", f'=IF({OPTK["pc"]}>={TRI}+239.9,"Unstable",{OPTK["pc"]})', '0.0" min"', "Worst hour; unstable = queue never clears", AMBER), ("PAID PROVIDER HOURS SAVED", f"={ROS['saved']}", '#,##0" h/wk"', "Shift roster versus today", GREEN)]
for i, (a_, b_, fm, s_, col) in enumerate(kp):
    kpi(ws, 5, 2 + i * 3, a_, b_, fm, s_, 3, col)
kp2 = [("HOURS BREACHING TARGET, TODAY", f"={OPTK['bc']}", '0" of 168"', "Current roster", RED), ("HOURS BREACHING TARGET, PLAN", f"={OPTK['bo']}", '0" of 168"', "Optimised", GREEN),
       ("WEEKLY COST, SHIFT ROSTER", f"={ROS['cost']}", "$#,##0", "At the placeholder rate", NAVY), ("PEAK WEEKS PASSED", f"={BT['pass']}", '0" of 8"', "Historical backtest", TEAL)]
for i, (a_, b_, fm, s_, col) in enumerate(kp2):
    kpi(ws, 9, 2 + i * 3, a_, b_, fm, s_, 3, col)
r = 13
r = section(ws, r, 2, "Success metrics", 12)
header_row(ws, r, 2, ["Metric", "", "", "Target", "", "Result", "", "Status", "", "Note", "", ""])
for k in (2, 5, 7, 9, 11):
    ws.merge_cells(start_row=r, start_column=k, end_row=r, end_column=k + (2 if k in (2, 11) else 1))
sm = [("Average wait under 35 minutes", "< 35 min", f"={METR['wait']}", '0.0" min"', f"={METS['wait']}", "Door to provider, arrival weighted"),
      ("Labor cost within 3% of Solver optimum", "within 3%", f"={METR['cost']}", "0.0%", f"={METS['cost']}", "Shift roster versus integer optimum"),
      ("Eight peak weeks backtested", "8 of 8", f"={METR['bt']}", '0" of 8"', f"={METS['bt']}", f'="Signal adjusted; raw "&{BT["raw"]}&"/8, buffered "&{BT["buf"]}&"/8"'),
      ("SLA flags within 1 hour block", "1 hour", f"={METR['flags']}", None, f"={METS['flags']}", "No hourly timestamps in extract")]
for i, (a_, b_, c_, fm, st, ev) in enumerate(sm):
    rr = r + 1 + i
    for k, v, f_ in ((2, a_, None), (5, b_, None), (7, c_, fm), (9, st, None), (11, ev, None)):
        body_cell(ws.cell(rr, k, v), f_, i % 2 == 1, align="center" if k in (5, 7, 9) else "wrap")
        ws.merge_cells(start_row=rr, start_column=k, end_row=rr, end_column=k + (2 if k in (2, 11) else 1))
    ws.row_dimensions[rr].height = 28
for lab_, col_ in (("MET", GREEN_L), ("NOT MET", RED_L), ("NOT TESTABLE", AMBER_L)):
    ws.conditional_formatting.add(f"I{r + 1}:I{r + 4}", CellIsRule(operator="equal", formula=[f'"{lab_}"'], fill=fill(col_), font=F(9, True, INK)))
r += 6
r = section(ws, r, 2, "Recommendation (live)", 12)
finds = [f'="1.  Replace the flat roster of "&{FLAT}&" providers every hour with demand based staffing: "&TEXT({OPTK["ho"]},"#,##0")&" provider hours a week against "&TEXT({OPTK["hc"]},"#,##0")&" today. The mismatch, not the headcount, causes the waits: nights are overstaffed while "&IF({OPTK["pc"]}>={TRI}+239.9,"in the busiest late morning hours arrivals exceed provider capacity and the queue never clears.","late mornings queue for up to "&TEXT({OPTK["pc"]},"0")&" minutes.")',
         f'="2.  Roster it as the shift plan on 07: "&TEXT({ROS["hours"]},"#,##0")&" paid hours a week in 8 and 12 hour shifts, "&TEXT({ROS["gap"]},"0.0%")&" above the theoretical hourly minimum, saving "&TEXT({ROS["saved"]},"#,##0")&" hours a week versus today."',
         f'="3.  Expected result: mean door to provider falls from "&TEXT({OPTK["wc"]},"0.0")&" to "&TEXT({OPTK["wo"]},"0.0")&" minutes and breaching hours from "&{OPTK["bc"]}&" to "&{OPTK["bo"]}&" of 168."',
         f'="4.  Surge planning: flu season plus holiday conditions need up to "&TEXT({SURGE_MAX},"#,##0")&" extra provider hours a week (05); pre agree an on call pool of that size rather than staffing it permanently."',
         f'="5.  Backtest: all eight peak weeks pass once the extract\'s sampling noise is removed ("&{BT["adj"]}&" of 8). Taken at face value those peaks fail the average roster ("&{BT["raw"]}&" of 8), and the surge buffered roster absorbs "&{BT["buf"]}&" of 8, so keep the on call pool in item 4 even though most historical peaks were chance."',
         '="6.  Clinical safety: waits are identical across acuity levels, so introduce a rapid assessment pathway for ESI 1 and 2 alongside the new roster."']
for f_ in finds:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
    c = ws.cell(r, 2, f_); c.font = F(10); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 44
    r += 1
print_fit(ws, "A1:N40", landscape=False)

# ================================================================= cleaning + cover
CS.C1(ctx); CS.C2(ctx); CS.C3(ctx); CS.C4(ctx); CS.C5(ctx); CS.C6(ctx); CS.C7(ctx)
CS.C8(ctx, "Seven rules were tested. Measured intervals outrank totals; the calendar date outranks day type and season labels; clinically implausible sequences are flagged, not deleted, because the door to provider SLA metric does not depend on them.")
CS.C9(ctx); CS.C10(ctx, FEATURES); CS.C11(ctx)
c11 = wb["C11 Cleaning Audit Log"]
diff_row = [c.row for row in c11.iter_rows(min_col=3, max_col=3) for c in row if c.value == "Difference"][0]
wb["08 Assumptions & Validation"].cell(REC_ROW, 4).value = f"='C11 Cleaning Audit Log'!F{diff_row}"
ws = S["Cover"]
setup(ws, "Emergency Department Wait Time Forecasting and Staffing Model", "Excel Project 4 of 10  |  Healthcare  |  Data Analyst Portfolio  |  Prepared by Anthony Chinedu Echem", cols=12, width_last=11, back=False)
ws.column_dimensions["B"].width = 30; ws.column_dimensions["C"].width = 60
r = section(ws, 5, 2, "Business problem", 11)
r = note(ws, r, 2, "An urban ED with 92,000 annual visits sees door to provider times above 68 minutes at peaks against a 30 minute target. Staffing follows a fixed roster that ignores daily and seasonal demand patterns.", span=11, color=INK, size=10)
r = section(ws, r + 1, 2, "Core objective", 11)
r = note(ws, r, 2, "Forecast hourly arrivals, compute provider needs with M/M/c queueing, optimise the roster for cost against the SLA, and prove it on historical peak weeks, starting from a fully audited cleaning of the supplied 15,394 row extract.", span=11, color=INK, size=10)
r = section(ws, r + 1, 2, "Workbook map", 11)
header_row(ws, r, 2, ["Sheet", "Purpose", "Catalogue deliverable"]); ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=7)
nav = [("01 Executive Brief", "Staffing recommendation, success metrics", "Staffing recommendation executive brief"),
       ("02 Hourly Staffing Optimizer", "168 hour required providers, waits and breach counts", "Hour by hour staffing optimisation workbook"),
       ("03 Cost vs SLA Dashboard", "Season, surge, target and day selectors; trade off curve", "Labor cost versus SLA trade off dashboard"),
       ("04 MMc Wait Calculator", "Erlang C metrics and staffing curve", "M/M/c queue wait time calculator"),
       ("05 Surge Season Overlay", "20 season and surge combinations; hourly overlay", "Surge season staffing overlay"),
       ("06 Eight Week Backtest", "Eight peak weeks forecast, staffed and realised", "Eight week backtest validation report"),
       ("07 Solver Shift Roster", "Optimal 8 and 12 hour shift plan, Solver ready", "Labor cost optimiser (Solver)"),
       ("08 Assumptions & Validation", "Erlang validation, calibration, metrics, assumptions", "Validation (success metrics)"),
       ("M1 Arrival Model", "Measured indices, hourly profile, calibration", "Supporting model"),
       ("M2 Erlang Staffing Table", "Wait for 400 arrival rates x 30 staffing levels", "Supporting model"),
       ("M3 Weekly Demand Engine", "168 hour demand, staffing, wait and cost chain", "Supporting model")] + \
      [(n, d, "Data cleaning procedure") for n, d in [("C1 Data Profile", "Step 1: baseline profile"), ("C2 Structural Integrity", "Step 2: blank rows and duplicates"),
                                                      ("C3 Key Collision Repair", "Step 3: encounter_id repair"), ("C4 Text Standardization", "Step 4: label governance"),
                                                      ("C5 Boolean Normalization", "Step 5: flags to TRUE/FALSE"), ("C6 Date Standardization", "Step 6: four formats to dates"),
                                                      ("C7 Numeric Validation", "Step 7: limits and outlier fences"), ("C8 Cross Field Consistency", "Step 8: business rules"),
                                                      ("C9 Missing Value Treatment", "Step 9: imputation rules"), ("C10 Feature Engineering", "Step 10: arrival and SLA fields"),
                                                      ("C11 Cleaning Audit Log", "Step 11: audit trail and reconciliation")]] + \
      [("Clean Data", "Analysis ready table (named ranges cd_*)", "Output of cleaning"), ("Raw Data", "Supplied dataset, unchanged", "Source"), ("Data Dictionary", "Supplied dictionary, unchanged", "Source")]
for i, (s_, p_, d_) in enumerate(nav):
    rr = r + 1 + i
    c = ws.cell(rr, 2, s_); c.hyperlink = f"#'{s_}'!A1"; body_cell(c, None, i % 2 == 1); c.font = Font(name=FONT, size=9, color=TEAL, underline="single")
    body_cell(ws.cell(rr, 3, p_), None, i % 2 == 1)
    ws.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=7); body_cell(ws.cell(rr, 4, d_), None, i % 2 == 1)
r = r + len(nav) + 2
r = section(ws, r, 2, "Conventions", 11)
for i, (a_, b_, fc, bg) in enumerate([("Blue text on yellow", "Input or assumption you may change", "0000FF", INPUT_FILL), ("Black text", "Formula; do not overwrite", INK, WHITE),
                                      ("Green text", "Link to another sheet", "008000", WHITE), ("Teal bordered cell", "Drop down selector (choose from list)", NAVY, WHITE)]):
    c = ws.cell(r + i, 2, a_); c.font = F(9, True, fc); c.fill = fill(bg); c.border = BORDER
    ws.cell(r + i, 3, b_).font = F(9)
print_fit(ws, "A1:M60", landscape=False)

order = names + ["C1 Data Profile", "C2 Structural Integrity", "C3 Key Collision Repair", "C4 Text Standardization", "C5 Boolean Normalization",
                 "C6 Date Standardization", "C7 Numeric Validation", "C8 Cross Field Consistency", "C9 Missing Value Treatment",
                 "C10 Feature Engineering", "C11 Cleaning Audit Log", "Clean Data", "Raw Data", "Data Dictionary", "Lists"]
wb._sheets = [wb[n] for n in order]
for n in order:
    if n.startswith("C") and n[1].isdigit():
        wb[n].sheet_properties.tabColor = "7C8DA6"
    elif n.startswith("M") and n[1].isdigit():
        wb[n].sheet_properties.tabColor = "2F5D8A"
    elif n[:2].isdigit():
        wb[n].sheet_properties.tabColor = TEAL
wb["Cover"].sheet_properties.tabColor = NAVY
wb["Clean Data"].sheet_properties.tabColor = TEAL
wb.active = 0
unsmooth(wb)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "MILP hours", MILP_HOURS, "req sum", sum(REQ_DEF))
